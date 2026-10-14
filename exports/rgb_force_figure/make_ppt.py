from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
out=Path(__file__).resolve().parent
r=Presentation(); r.slide_width=Inches(16); r.slide_height=Inches(9)
s=r.slides.add_slide(r.slide_layouts[6])
C={'ink':'233247','sub':'58697B','rgb':'FBE6DF','ft':'DEEFF6','q':'EBE3F8','green':'E3F1E7','line':'708296','panel':'F3F6FA','white':'FFFFFF'}
def text(x,y,w,h,txt,size=18,bold=False,color='ink',align=PP_ALIGN.CENTER):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.035);tf.margin_top=tf.margin_bottom=Inches(.025)
 tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for i,t in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.alignment=align
  p.font.name='Noto Sans CJK KR';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color))
 return sh
def box(x,y,w,h,txt,fill='panel',size=18,bold=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C[fill]);sh.line.color.rgb=RGBColor.from_string('C2CDD9');sh.line.width=Pt(1)
 sh.adjustments[0]=.1
 text(x+.05,y+.025,w-.1,h-.05,txt,size,bold)
 return sh
def arrow(x,y,X,Y,color='line',dash=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(X),Inches(Y));sh.line.color.rgb=RGBColor.from_string(C[color]);sh.line.width=Pt(1.7)
 end=OxmlElement('a:tailEnd');end.set('type','triangle');sh.line._get_or_add_ln().append(end)
 return sh
text(.4,.3,15.2,.65,'사진과 힘의 정보가 두 Transformer 층을 거쳐 Query에 모이는 과정',26,True,align=PP_ALIGN.LEFT)
text(.43,1.03,15,.4,'입력 5개 → 1층 출력 5개 → 2층 출력 5개 → 최종 Query 하나로 상태 예측',18,color='sub',align=PP_ALIGN.LEFT)
for x,w,t in [(.35,2.25,'입력 · 특징 추출'),(2.95,1.8,'입력 토큰'),(6.7,1.8,'1층 출력'),(10.45,1.8,'2층 출력'),(12.9,2.7,'상태 예측')]:
 text(x,1.77,w,.45,t,19,True)
ys=[2.6,3.4,4.2,5.0,5.8]
labels=['이전 사진','현재 사진','왼쪽 힘','오른쪽 힘','Query']
colors=['rgb','rgb','ft','ft','q']
box(.35,2.6,2.25,1.4,'사진 2장 → CLIP\n+ Linear\n0.05초 간격 · CLIP 고정','rgb',15)
box(.35,4.2,2.25,.6,'왼쪽 F/T → CNN','ft',16)
box(.35,5.0,2.25,.6,'오른쪽 F/T → CNN','ft',16)
box(.35,5.8,2.25,.6,'학습 가능한 Query','q',15)
for y in ys: arrow(2.63,y+.3,2.91,y+.3)
for x,stage in [(2.95,0),(6.7,1),(10.45,2)]:
 for y,label,col in zip(ys,labels,colors):
  suffix='' if stage==0 else (' ①' if stage==1 else ' ②')
  box(x,y,1.8,.6,label+suffix,col,16,True)
 text(x,6.57,1.8,.63,'5개 토큰\n각각 128개 숫자',14,color='sub')
# Each layer consumes all five tokens and emits all five updated tokens.
for x,title,left,right in [(5.12,'트랜스포머\n1층',4.75,6.7),(8.87,'트랜스포머\n2층',8.5,10.45)]:
 box(x,2.43,1.23,4.18,'','panel')
 text(x+.03,3.87,1.17,1.05,title,14,True)
 text(x+.04,5.55,1.15,.72,'정보 비교\n및 결합',12,color='sub')
 for y in ys:
  arrow(left+.035,y+.3,x-.04,y+.3)
  arrow(x+1.26,y+.3,right-.045,y+.3)
# Only the purple query output is connected to the classifier.
arrow(12.29,6.1,12.86,6.1,color='ink')
box(12.9,5.68,2.7,.84,'최종 Query만 선택\n128개 숫자','q',17,True)
arrow(14.25,5.64,14.25,5.25)
box(12.9,4.38,2.7,.83,'Linear: 128 → 4\n네 상태의 점수','panel',17)
arrow(14.25,4.34,14.25,3.97)
box(12.9,3.38,2.7,.55,'Softmax → 확률','green',17)
arrow(14.25,3.34,14.25,3.03)
box(12.9,2.42,2.7,.57,'접근 / 회전 / 복구 / 오류','green',14,True)
text(12.87,6.62,2.75,.64,'다른 4개 출력 토큰은\n최종 분류층에 넣지 않음',13,color='sub')
text(.35,6.62,2.25,.65,'F/T: 약 0.4초 기록\n원본·평균·변화량 사용',13,color='sub')
box(.35,7.5,15.25,.62,'①·②는 각 층에서 갱신된 특징을 뜻함  |  사진과 힘의 정보가 반영된 최종 Query로만 분류','panel',17)
text(.35,8.25,15.25,.39,'학습: 현재 시점의 정답 라벨과 예측을 비교하여 오차 최소화 · Query는 정답 라벨이 아님',15,color='sub')
text(.35,8.7,15.25,.2,'시간 기준: RGB 60 fps / F/T 100 Hz. 종류·위치 임베딩과 Transformer 내부 연산은 생략. 각 층에서의 정보 교환은 causal attention mask를 따름.',10,color='sub')
s.notes_slide.notes_text_frame.text='입력 토큰 다섯 개가 1층에 들어가 갱신된 토큰 다섯 개로 나옵니다. 이 다섯 개가 모두 2층에 들어가 다시 갱신된 다섯 개로 나옵니다. 마지막 다섯 개 중 query 출력만 분류층에 사용합니다. 각 토큰은 128차원입니다. 그림의 ①, ②는 갱신 단계를 구분하는 표시입니다.'
p=out/'rgb_force_context_simple.pptx';r.save(p)
rr=Presentation(p);assert len(rr.slides)==1
print(p)
print('Editable shapes:',len(rr.slides[0].shapes))
